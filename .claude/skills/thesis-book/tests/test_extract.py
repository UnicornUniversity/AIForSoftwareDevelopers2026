from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

import docx

from bookbuild.extract import Extractor, extract_docx
from bookbuild.manifest import load_manifest

from tests.helpers import PROJECT, make_fixture


class ExtractTests(unittest.TestCase):
    def test_docx_body_and_metadata(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            fixture = make_fixture(Path(directory))
            blocks = extract_docx(fixture["corpus"] / "doc01-fixture.docx")
        kinds = [block["kind"] for block in blocks]
        texts = [block["text"] for block in blocks]
        self.assertEqual(kinds[0], "meta")
        self.assertIn("Executive summary", texts)
        self.assertTrue(
            any("alpha line" in text and kind == "paragraph" for text, kind in zip(texts, kinds))
        )
        self.assertTrue(any("—" in text for text in texts))

    def test_cache_reuse_and_invalidation(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fixture = make_fixture(root)
            manifest = load_manifest(fixture["manifest"])
            entry = manifest.lookup("doc01")
            extractor = Extractor(root / "cache")
            first = extractor.load_or_extract(entry, entry.path(fixture["corpus"]))
            cache_file = root / "cache" / "doc01.json"
            self.assertTrue(cache_file.is_file())

            second = extractor.load_or_extract(entry, entry.path(fixture["corpus"]))
            self.assertEqual(first.blocks, second.blocks)

            document = docx.Document(str(entry.path(fixture["corpus"])))
            document.add_paragraph("A newly appended sentence changes the hash.")
            document.save(str(entry.path(fixture["corpus"])))
            third = extractor.load_or_extract(entry, entry.path(fixture["corpus"]))
            self.assertNotEqual(first.sha256, third.sha256)
            self.assertTrue(any("newly appended" in text for text in third.paragraphs))

    def test_real_corpus_strips_running_headers(self) -> None:
        project = PROJECT
        manifest_path = project / "coprus" / "manifest.json"
        if not manifest_path.is_file():
            self.skipTest("the corpus is not present")
        manifest = load_manifest(manifest_path)
        entry = manifest.lookup("doc02")
        with tempfile.TemporaryDirectory() as directory:
            extractor = Extractor(Path(directory))
            source = extractor.load_or_extract(entry, entry.path(project / "coprus"))
        joined = "\n".join(source.paragraphs)
        self.assertIn("Dusk Chorus Behavior", joined)
        self.assertIn("1. Introduction", joined)
        self.assertNotIn("https://doi.org", joined)
        self.assertNotIn("Australasian Journal of Macropod Ecology 41(2) (2023)", joined)


if __name__ == "__main__":
    unittest.main()
