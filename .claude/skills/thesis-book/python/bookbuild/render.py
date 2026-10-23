"""Render the plan, front matter and chapter files into one Word document.

Every render starts from a fresh ``Document`` and writes the whole book, so the
output is a pure function of the source files.  The binary is never patched.
Word resolves the table of contents and page fields when the document opens
(``w:updateFields`` asks it to).
"""

from __future__ import annotations

import os
import re
from pathlib import Path

from .errors import ContentError, RenderError
from .manifest import DocumentEntry, Manifest
from .markdown import Block, Chapter
from .plan import Plan
from .settings import Settings
from .typography import apply as apply_typography

_PAGE_WIDTH_CM = {"A4": 21.0, "Letter": 21.59}
_PAGE_HEIGHT_CM = {"A4": 29.7, "Letter": 27.94}


def _qn(tag: str):
    from docx.oxml.ns import qn

    return qn(tag)


def _set_rfonts(rpr, font: str) -> None:
    from docx.oxml import OxmlElement

    rfonts = rpr.find(_qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(_qn(attribute), font)
    for attribute in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        key = _qn(attribute)
        if key in rfonts.attrib:
            del rfonts.attrib[key]


def _clear_theme_colour(rpr) -> None:
    colour = rpr.find(_qn("w:color"))
    if colour is None:
        return
    for attribute in ("w:themeColor", "w:themeShade", "w:themeTint"):
        key = _qn(attribute)
        if key in colour.attrib:
            del colour.attrib[key]


def _configure_style(
    style,
    *,
    font: str,
    size: float,
    bold: bool = False,
    italic: bool = False,
    colour: str | None = None,
    alignment=None,
    line_spacing: float | None = None,
    space_before: float | None = None,
    space_after: float | None = None,
    first_line_indent_cm: float | None = None,
    left_indent_cm: float | None = None,
    right_indent_cm: float | None = None,
    keep_with_next: bool = False,
) -> None:
    from docx.shared import Cm, Pt, RGBColor

    style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = italic
    if colour is not None:
        style.font.color.rgb = RGBColor.from_string(colour)
    rpr = style.element.get_or_add_rPr()
    _set_rfonts(rpr, font)
    _clear_theme_colour(rpr)

    paragraph = style.paragraph_format
    if alignment is not None:
        paragraph.alignment = alignment
    if line_spacing is not None:
        paragraph.line_spacing = line_spacing
    if space_before is not None:
        paragraph.space_before = Pt(space_before)
    if space_after is not None:
        paragraph.space_after = Pt(space_after)
    if first_line_indent_cm is not None:
        paragraph.first_line_indent = Cm(first_line_indent_cm)
    if left_indent_cm is not None:
        paragraph.left_indent = Cm(left_indent_cm)
    if right_indent_cm is not None:
        paragraph.right_indent = Cm(right_indent_cm)
    paragraph.keep_with_next = keep_with_next


def _build_styles(document, settings: Settings) -> dict[str, object]:
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    typography = settings.raw["typography"]
    body = typography["body"]
    headings = typography["headings"]
    front = typography["front_titles"]
    quote = typography["quote"]
    provenance = typography["provenance"]

    normal = document.styles["Normal"]
    _configure_style(
        normal,
        font=body["font"],
        size=body["size_pt"],
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=body["line_spacing"],
        space_before=body["space_before_pt"],
        space_after=body["space_after_pt"],
        first_line_indent_cm=body["first_line_indent_cm"],
    )

    for level, key, before, after in (
        (1, "h1_size_pt", "h1_space_before_pt", "h1_space_after_pt"),
        (2, "h2_size_pt", "h2_space_before_pt", "h2_space_after_pt"),
    ):
        style = document.styles[f"Heading {level}"]
        _configure_style(
            style,
            font=headings["font"],
            size=headings[key],
            bold=headings["bold"],
            colour=headings["color"],
            alignment=WD_ALIGN_PARAGRAPH.LEFT,
            line_spacing=1.0,
            space_before=headings[before],
            space_after=headings[after],
            first_line_indent_cm=0,
            keep_with_next=True,
        )

    def add(name: str, **kwargs):
        style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style.base_style = normal
        _configure_style(style, **kwargs)
        return style

    add(
        "Book Front Title",
        font=front["font"],
        size=front["size_pt"],
        bold=front["bold"],
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=18,
        keep_with_next=True,
    )
    add(
        "Book Title Page",
        font=front["font"],
        size=24,
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=12,
    )
    add(
        "Book Title Subtitle",
        font=front["font"],
        size=16,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=12,
    )
    add(
        "Book Title Meta",
        font=body["font"],
        size=12,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        space_before=0,
        space_after=6,
    )
    add(
        "Book Quote",
        font=body["font"],
        size=quote["size_pt"],
        italic=quote["italic"],
        alignment=WD_ALIGN_PARAGRAPH.JUSTIFY,
        line_spacing=1.0,
        space_before=6,
        space_after=0,
        first_line_indent_cm=0,
        left_indent_cm=quote["indent_cm"],
        right_indent_cm=quote["indent_cm"],
    )
    add(
        "Book Provenance",
        font=body["font"],
        size=provenance["size_pt"],
        italic=provenance["italic"],
        colour=provenance["color"],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=2,
        space_after=12,
        first_line_indent_cm=0,
        left_indent_cm=quote["indent_cm"],
    )
    add(
        "Book List",
        font=body["font"],
        size=body["size_pt"],
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=body["line_spacing"],
        space_before=0,
        space_after=0,
        first_line_indent_cm=0,
        left_indent_cm=1.0,
    )
    add(
        "Book Caption",
        font=body["font"],
        size=10,
        italic=True,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        space_before=6,
        space_after=4,
        first_line_indent_cm=0,
        keep_with_next=True,
    )
    add(
        "Book Reference",
        font=body["font"],
        size=10,
        alignment=WD_ALIGN_PARAGRAPH.LEFT,
        line_spacing=1.0,
        space_before=0,
        space_after=4,
        first_line_indent_cm=0,
        left_indent_cm=0.6,
    )
    return {}


def _add_field(paragraph, instruction: str, placeholder: str = "") -> None:
    from docx.oxml import OxmlElement

    begin = OxmlElement("w:fldChar")
    begin.set(_qn("w:fldCharType"), "begin")
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(_qn("xml:space"), "preserve")
    instruction_element.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(_qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(_qn("w:fldCharType"), "end")

    run = paragraph.add_run()
    run._r.append(begin)
    run._r.append(instruction_element)
    run._r.append(separate)
    run._r.append(text)
    run._r.append(end)


def _update_fields(document) -> None:
    from docx.oxml import OxmlElement

    settings = document.settings.element
    update = settings.find(_qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(_qn("w:val"), "true")


def _restart_page_numbering(section) -> None:
    from docx.oxml import OxmlElement

    sect_pr = section._sectPr
    existing = sect_pr.find(_qn("w:pgNumType"))
    if existing is not None:
        sect_pr.remove(existing)
    page_numbers = OxmlElement("w:pgNumType")
    page_numbers.set(_qn("w:start"), "1")
    columns = sect_pr.find(_qn("w:cols"))
    if columns is not None:
        columns.addprevious(page_numbers)
    else:
        sect_pr.append(page_numbers)


def _page_setup(section, settings: Settings) -> None:
    from docx.shared import Cm

    page = settings.raw["page"]
    width = _PAGE_WIDTH_CM[page["size"]]
    height = _PAGE_HEIGHT_CM[page["size"]]
    section.page_width = Cm(width)
    section.page_height = Cm(height)
    margins = page["margins_cm"]
    section.top_margin = Cm(margins["top"])
    section.bottom_margin = Cm(margins["bottom"])
    section.left_margin = Cm(margins["left"])
    section.right_margin = Cm(margins["right"])
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)


def _text(settings: Settings, value: str) -> str:
    return apply_typography(value, settings.typography)


def _add_title_page(document, settings: Settings, frontmatter: dict) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt

    book = settings.raw["book"]
    title = document.add_paragraph(style="Book Title Page")
    title.paragraph_format.space_before = Pt(170)
    title.add_run(_text(settings, book["title"]))

    if book.get("subtitle"):
        subtitle = document.add_paragraph(style="Book Title Subtitle")
        subtitle.add_run(_text(settings, book["subtitle"]))

    author = document.add_paragraph(style="Book Title Meta")
    author.paragraph_format.space_before = Pt(150)
    author.add_run(_text(settings, book["author"]))

    if book.get("institution"):
        institution = document.add_paragraph(style="Book Title Meta")
        institution.add_run(_text(settings, book["institution"]))

    date = document.add_paragraph(style="Book Title Meta")
    date.add_run(_text(settings, book["date"]))
    date.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _front_heading(document, settings: Settings, text: str) -> None:
    paragraph = document.add_paragraph(style="Book Front Title")
    paragraph.add_run(_text(settings, text))


def _add_text_page(
    document,
    settings: Settings,
    heading: str,
    text: str,
    *,
    centre: bool = False,
) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    _front_heading(document, settings, heading)
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = None
    if centre:
        paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = None
    paragraph.add_run(_text(settings, text))


def _add_keywords(document, settings: Settings, label: str, keywords: list[str]) -> None:
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = None
    bold = paragraph.add_run(f"{_text(settings, label)}: ")
    bold.bold = True
    paragraph.add_run(_text(settings, "; ".join(keywords)))


def _add_abstract(
    document, settings: Settings, frontmatter: dict, key: str, label: str, keywords_label: str
) -> None:
    section = frontmatter[key]
    _front_heading(document, settings, label)
    for paragraph_text in _split_paragraphs(section["text"]):
        paragraph = document.add_paragraph(style="Normal")
        paragraph.paragraph_format.first_line_indent = None
        paragraph.add_run(_text(settings, paragraph_text))
    _add_keywords(document, settings, keywords_label, section["keywords"])


def _split_paragraphs(text: str) -> list[str]:
    parts = [re.sub(r"\s+", " ", part).strip() for part in re.split(r"\n\s*\n", text)]
    return [part for part in parts if part]


def _add_table_of_contents(document, settings: Settings) -> None:
    toc = settings.get("front_matter.table_of_contents", {})
    _front_heading(document, settings, toc["title"])
    paragraph = document.add_paragraph(style="Normal")
    paragraph.paragraph_format.first_line_indent = None
    depth = int(toc["depth"])
    _add_field(
        paragraph,
        f'TOC \\o "1-{depth}" \\h \\z \\u',
        "Right-click and choose Update Field to build the contents.",
    )


def _add_front_matter(document, settings: Settings, frontmatter: dict) -> None:
    flags = settings.raw["front_matter"]
    started = False

    def begin_section() -> None:
        nonlocal started
        if started:
            document.add_page_break()
        started = True

    if flags["title_page"]:
        begin_section()
        _add_title_page(document, settings, frontmatter)

    dedication = frontmatter.get("dedication", "").strip()
    if flags["dedication"] and dedication:
        begin_section()
        _add_text_page(document, settings, "Dedication", dedication, centre=True)

    declaration = frontmatter.get("declaration", "").strip()
    if flags["declaration"] and declaration:
        begin_section()
        _add_text_page(document, settings, "Declaration", declaration)

    abstract = flags["abstract"]
    if abstract["enabled"]:
        begin_section()
        _add_abstract(
            document,
            settings,
            frontmatter,
            "abstract",
            abstract["title"],
            abstract["keywords_title"],
        )

    anotace = flags["anotace"]
    if anotace["enabled"]:
        begin_section()
        _add_abstract(
            document,
            settings,
            frontmatter,
            "anotace",
            anotace["title"],
            anotace["keywords_title"],
        )

    if flags["table_of_contents"]["enabled"]:
        begin_section()
        _add_table_of_contents(document, settings)


def _configure_body_section(document, settings: Settings, title: str):
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    body = document.add_section(WD_SECTION.NEW_PAGE)
    _page_setup(body, settings)

    header_settings = settings.raw["body"]["running_header"]
    body.header.is_linked_to_previous = False
    body.footer.is_linked_to_previous = False
    for paragraph in body.header.paragraphs:
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
    if header_settings["enabled"]:
        paragraph = body.header.paragraphs[0]
        paragraph.alignment = {
            "left": WD_ALIGN_PARAGRAPH.LEFT,
            "center": WD_ALIGN_PARAGRAPH.CENTER,
            "right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[header_settings["align"]]
        if header_settings["text"] == "chapter_title":
            _add_field(paragraph, 'STYLEREF "Heading 1" \\* MERGEFORMAT', "")
        else:
            paragraph.add_run(_text(settings, title))

    for paragraph in body.footer.paragraphs:
        for run in list(paragraph.runs):
            run._r.getparent().remove(run._r)
    numbers = settings.raw["body"]["page_numbers"]
    if numbers["enabled"]:
        paragraph = body.footer.paragraphs[0]
        paragraph.alignment = {
            "bottom_center": WD_ALIGN_PARAGRAPH.CENTER,
            "bottom_right": WD_ALIGN_PARAGRAPH.RIGHT,
        }[numbers["position"]]
        _add_field(paragraph, "PAGE", "1")
    if numbers["start_at_introduction"]:
        _restart_page_numbering(body)
    return body


def _citation_text(
    settings: Settings, manifest: Manifest, doc_id: str
) -> str:
    entry = manifest.lookup(doc_id)
    if settings.get("citations.in_text", "author_date") == "id":
        return f"({doc_id})"
    return entry.citation


def _add_inline_runs(
    paragraph, block: Block, settings: Settings, manifest: Manifest
) -> None:
    for span in block.inlines:
        if span.citation:
            doc_id = span.text[1:-1]
            text = _citation_text(settings, manifest, doc_id)
        else:
            text = _text(settings, span.text)
        run = paragraph.add_run(text)
        if span.bold:
            run.bold = True
        if span.italic:
            run.italic = True


def _provenance(settings: Settings, entry: DocumentEntry) -> str:
    year = entry.year if entry.year is not None else "n.d."
    return _text(
        settings,
        f"{entry.id} · {entry.author_label} ({year}) · {entry.title} "
        f"({entry.format_label}, {entry.word_count} words)",
    )


def _add_chapter(
    document,
    settings: Settings,
    manifest: Manifest,
    chapter: Chapter,
    *,
    first: bool,
) -> None:
    heading = document.add_paragraph(style="Heading 1")
    if settings.get("body.chapter_start_new_page", True) and not first:
        heading.paragraph_format.page_break_before = True
    heading.add_run(_text(settings, f"{chapter.number}  {chapter.title}"))

    section_number = 0
    for block in chapter.blocks:
        if block.kind == "heading":
            section_number += 1
            subheading = document.add_paragraph(style="Heading 2")
            subheading.add_run(
                _text(settings, f"{chapter.number}.{section_number}  {block.text}")
            )
        elif block.kind == "paragraph":
            paragraph = document.add_paragraph(style="Normal")
            _add_inline_runs(paragraph, block, settings, manifest)
        elif block.kind == "list_item":
            paragraph = document.add_paragraph(style="Book List")
            paragraph.add_run("•\t")
            _add_inline_runs(paragraph, block, settings, manifest)
        elif block.kind == "quote":
            quote = document.add_paragraph(style="Book Quote")
            quote.add_run(_text(settings, block.text))
            entry = manifest.lookup(block.source or "")
            provenance = document.add_paragraph(style="Book Provenance")
            provenance.add_run(_provenance(settings, entry))
        else:  # pragma: no cover - the parser only emits the kinds above
            raise RenderError(f"unsupported block kind {block.kind!r}")


def _authors_harvard(authors: tuple[str, ...]) -> str:
    if not authors:
        return ""
    formatted: list[str] = []
    for name in authors:
        parts = name.split()
        if len(parts) == 1:
            formatted.append(parts[0])
        else:
            formatted.append(f"{parts[-1]}, {' '.join(p[0] + '.' for p in parts[:-1])}")
    if len(formatted) == 1:
        return formatted[0]
    if len(formatted) == 2:
        return f"{formatted[0]}, & {formatted[1]}"
    return ", ".join(formatted[:-1]) + f", & {formatted[-1]}"


def _reference_entry(entry: DocumentEntry, style: str) -> str:
    authors = _authors_harvard(entry.authors) or entry.id
    year = entry.year if entry.year is not None else "n.d."
    title = entry.title.rstrip(".")
    venue = entry.venue.rstrip(".")
    parts = [f"{authors} ({year}). {title}."]
    if venue:
        parts.append(f" {venue}.")
    return "".join(parts)


def _add_references(
    document,
    settings: Settings,
    manifest: Manifest,
    used_ids: list[str],
) -> None:
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.shared import Cm, Pt

    reference_settings = settings.raw["body"]["references"]
    ids = list(used_ids) if reference_settings["scope"] == "used" else [
        entry.id for entry in manifest.documents
    ]
    entries = [manifest.lookup(doc_id) for doc_id in ids]
    if reference_settings["sort"] == "title":
        entries.sort(key=lambda entry: entry.title.casefold())
    else:
        entries.sort(key=lambda entry: entry.id)

    heading = document.add_paragraph(style="Heading 1")
    heading.paragraph_format.page_break_before = True
    heading.add_run(_text(settings, reference_settings["title"]))

    captions = settings.raw["body"]["captions"]
    caption = "Table 1: Source documents cited in this thesis."
    if captions["table"] == "above":
        paragraph = document.add_paragraph(style="Book Caption")
        paragraph.add_run(_text(settings, caption))

    table = document.add_table(rows=1, cols=2)
    table.style = document.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = (Cm(2.2), Cm(12.8))
    for column, width in zip(table.columns, widths):
        column.width = width
    header_cells = table.rows[0].cells
    for index, label in enumerate(("ID", "Reference")):
        header_cells[index].width = widths[index]
        paragraph = header_cells[index].paragraphs[0]
        run = paragraph.add_run(label)
        run.bold = True
        run.font.size = Pt(10)
    for entry in entries:
        cells = table.add_row().cells
        cells[0].width, cells[1].width = widths
        first = cells[0].paragraphs[0].add_run(entry.id)
        first.font.size = Pt(10)
        second = cells[1].paragraphs[0].add_run(
            _text(settings, _reference_entry(entry, reference_settings["style"]))
        )
        second.font.size = Pt(10)

    # Repeat the header row when the table crosses a page.
    from docx.oxml import OxmlElement

    row_properties = table.rows[0]._tr.get_or_add_trPr()
    header_flag = OxmlElement("w:tblHeader")
    header_flag.set(_qn("w:val"), "true")
    row_properties.append(header_flag)

    if captions["table"] == "below":
        paragraph = document.add_paragraph(style="Book Caption")
        paragraph.add_run(_text(settings, caption))


def _used_ids(chapters: list[Chapter]) -> list[str]:
    used: set[str] = set()
    for chapter in chapters:
        for block in chapter.blocks:
            used.update(block.citations)
            if block.kind == "quote" and block.source:
                used.add(block.source)
    return sorted(used)


def render_book(
    settings: Settings,
    plan: Plan,
    manifest: Manifest,
    chapters: list[Chapter],
    frontmatter: dict,
) -> Path:
    """Build the whole document and return the written path."""

    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise RenderError(
            "python-docx is not installed; run: python -m pip install --user -r "
            f"{Path(__file__).resolve().parents[2] / 'requirements.txt'}"
        ) from exc

    document = Document()
    _page_setup(document.sections[0], settings)
    _build_styles(document, settings)
    _update_fields(document)

    _add_front_matter(document, settings, frontmatter)
    _configure_body_section(document, settings, plan.title)

    for index, chapter in enumerate(chapters):
        _add_chapter(
            document,
            settings,
            manifest,
            chapter,
            first=index == 0,
        )

    if settings.get("body.references.enabled", True):
        _add_references(document, settings, manifest, _used_ids(chapters))

    properties = document.core_properties
    properties.title = settings.raw["book"]["title"]
    properties.subject = settings.raw["book"].get("subtitle", "")
    properties.author = settings.raw["book"]["author"]
    properties.language = settings.raw["book"]["language"]

    output = settings.output_path
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(output.stem + ".tmp.docx")
    try:
        document.save(str(temporary))
        os.replace(temporary, output)
    except OSError as exc:
        raise RenderError(f"cannot write the document: {exc}", file=output) from exc
    finally:
        if temporary.exists():
            try:
                temporary.unlink()
            except OSError:
                pass
    return output
